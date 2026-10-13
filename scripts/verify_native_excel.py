"""Read, never author, an Excel-saved workbook and compare its cached calculations."""

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

from shipping_equity.data import load_company
from shipping_equity.valuation import reverse_valuation, value_company

parser = argparse.ArgumentParser(description=__doc__)
parser.add_argument("workbook", type=Path)
parser.add_argument("--company", type=Path, default=Path("data/companies/SBLK"))
parser.add_argument("--case", choices=["bear", "base", "bull"], required=True)
parser.add_argument("--output", type=Path, required=True)
args = parser.parse_args()
company = load_company(args.company)
wb = load_workbook(args.workbook, data_only=True)
formula_wb = load_workbook(args.workbook, data_only=False)
assert wb["Assumptions"]["D5"].value == {"bear": 1, "base": 2, "bull": 3}[args.case]
assert wb["Assumptions"]["D6"].value == 1 and wb["Assumptions"]["D7"].value == 0
result = value_company(company, args.case)
controls = []


def compare(label, excel_values, python_values, tolerance=0.0001):
    left, right = np.array(excel_values, dtype=float), np.array(python_values, dtype=float)
    assert left.shape == right.shape, label
    error = float(np.max(np.abs(left - right)))
    controls.append(
        {
            "control": label,
            "values": int(left.size),
            "max_abs_difference": error,
            "pass": bool(np.isfinite(error) and error <= tolerance),
        }
    )


for cell, key in {
    "D10": "price",
    "D11": "nav_per_share",
    "D12": "price_nav",
    "D16": "enterprise_value",
    "D18": "equity_dcf",
    "D19": "dcf_per_share",
    "D23": "next_year_revenue",
    "D24": "next_year_ufcf",
    "D25": "next_year_equity_cash",
}.items():
    compare(key, wb["Summary"][cell].value, result.summary[key])
layout_path = args.workbook.parent / "excel-calculation-evidence.json"
if not layout_path.exists():
    layout_path = (
        Path("outputs")
        / f"{company.metadata['ticker'].lower()}-{company.metadata['valuation_date']}"
        / "excel-calculation-evidence.json"
    )
layout = json.loads(layout_path.read_text())["layout"]
for sheet, table, keys in [
    ("Cash Flow", result.forecast, layout["forecast"]),
    ("Debt", result.financing, layout["debt"]),
]:
    for key, row in keys.items():
        compare(
            f"{sheet} / {key}",
            [wb[sheet].cell(row, col).value for col in range(5, 5 + len(table))],
            table[key].to_numpy(),
        )
for row, adjustment in enumerate([-0.02, -0.01, 0, 0.01, 0.02], start=9):
    for col, multiplier in enumerate([0.8, 0.9, 1, 1.1, 1.2], start=13):
        expected = value_company(
            company, args.case, rate_multiplier=multiplier, discount_adjustment=adjustment
        ).summary["dcf_per_share"]
        compare(f"Native data table {row},{col}", wb["Assumptions"].cell(row, col).value, expected)
root = reverse_valuation(company, args.case)
compare("Reverse valuation", wb["Summary"]["D30"].value, root["multiplier"], 1e-8)
errors = [
    f"{sheet.title}!{cell.coordinate}: {cell.value}"
    for sheet in wb
    for row in sheet
    for cell in row
    if cell.data_type == "e"
]
assert not errors, errors[:10]
assert all(row["pass"] for row in controls), controls
assert getattr(formula_wb["Assumptions"]["M9"].value, "t", None) == "dataTable"
record = {
    "status": "pass",
    "case": args.case,
    "engine": "Microsoft Excel for Mac: saved after live case-selector edit",
    "input_sha256": company.fingerprint,
    "workbook_sha256": hashlib.sha256(args.workbook.read_bytes()).hexdigest(),
    "values_compared": sum(row["values"] for row in controls),
    "formula_errors": errors,
    "native_two_input_table": "25 recalculated cells matched full Python valuations",
    "controls": controls,
}
args.output.write_text(json.dumps(record, indent=2) + "\n")
print(f"Native {args.case}: {record['values_compared']} values passed; no formula errors")
